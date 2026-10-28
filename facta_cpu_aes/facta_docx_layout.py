"""Render the shared manuscript into the official Facta Word template.

The SJEE module supplies only LaTeX/BibTeX parsing and native equation conversion.
All Word layout, figures, title matter and validation here are Facta-specific.
"""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
import subprocess
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent / 'sjee'))
import fill_sjee_template_docx as shared

TEMPLATE = ROOT / 'Facta_Template.docx'
OUTPUT = ROOT / 'main_facta.docx'
BUILD = ROOT / 'build/docx'
FIGURES = ROOT / 'figures/submission'
WIDTH_CM = 12.7


def fingerprint(path):
    return sha256(path.read_bytes()).hexdigest()


def math(tex):
    node = shared.omml(tex)
    for run in node.iter(qn('m:r')):
        props = run.find(qn('w:rPr'))
        # Word scales sub/superscripts automatically; forcing 6 pt here would
        # scale them a second time and make indices unreadably small.
        props.find(qn('w:sz')).set(qn('w:val'), '20')
        fonts = OxmlElement('w:rFonts')
        for name in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            fonts.set(qn('w:' + name), 'Times New Roman')
        props.append(fonts)
    return node


def write_runs(paragraph, runs):
    for item in runs:
        if 'math' in item:
            paragraph._p.append(math(item['math']))
        else:
            run = paragraph.add_run(item['text'])
            if 'bold' in item:
                run.bold = item['bold']
            if 'italic' in item:
                run.italic = item['italic']
            # Code identifiers use the journal's body font, not Courier.
            run.font.name = 'Times New Roman'


def hyperlink(paragraph, label, url):
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


def replace_footnote(document, data):
    part = next(p for p in document.part.package.parts if str(p.partname) == '/word/footnotes.xml')
    root = parse_xml(part.blob)
    note = next(n for n in root if n.get(qn('w:id')) == '1')
    for node in list(note):
        note.remove(node)
    for text in ['Received [date]', 'Corresponding author: ' + data['author'],
                 data['institution'], '(E-mail: ' + data['email'] + ')']:
        p = OxmlElement('w:p')
        pp = OxmlElement('w:pPr')
        style = OxmlElement('w:pStyle')
        style.set(qn('w:val'), 'FootnoteText')
        pp.append(style)
        p.append(pp)
        run = OxmlElement('w:r')
        node = OxmlElement('w:t')
        node.text = text
        run.append(node)
        p.append(run)
        note.append(p)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def figure_assets(relative):
    source = ROOT / relative
    assert source.suffix == '.pdf'
    target = FIGURES / source.stem
    # 1500 pixels across 12.7 cm = 300 pixels/inch at the inserted width.
    raster = BUILD / (source.stem + '_300dpi')
    subprocess.run(['pdftocairo', '-png', '-singlefile', '-r', '300',
                    '-scale-to-x', '1500', '-scale-to-y', '-1',
                    str(source), str(raster)], check=True, capture_output=True)
    with Image.open(raster.with_suffix('.png')) as image:
        rgb = image.convert('RGB')
        rgb.save(target.with_suffix('.jpg'), quality=100, subsampling=0, dpi=(300, 300))
        rgb.save(target.with_suffix('.tif'), compression='tiff_lzw', dpi=(300, 300))
    return target.with_suffix('.jpg')


def reference_runs(reference, entry):
    def value(key):
        return shared.plain(shared.parse(entry.get(key, '')))
    author = reference['text'].split(': ', 1)[0]
    runs = [dict(text=author + ', ')]
    report = entry['ENTRYTYPE'] == 'techreport'
    runs.append(dict(text=value('title') if report else '\u201c' + value('title') + '\u201d', italic=report))
    for key in ('journal', 'booktitle'):
        if value(key):
            runs.extend([dict(text=', '), dict(text=value(key), italic=True)])
    for key, prefix in [('institution', ''), ('type', ''), ('number', ''),
                        ('volume', 'vol. '), ('pages', 'pp. '), ('year', '')]:
        if value(key):
            runs.append(dict(text=', ' + prefix + value(key)))
    if value('doi'):
        runs.append(dict(text=', doi: ' + value('doi')))
    elif value('url'):
        runs.append(dict(text='. Available: ' + value('url')))
    runs.append(dict(text='.'))
    return runs


def render(data, temporary):
    document = Document(TEMPLATE)
    body = document._element.body
    table_properties = deepcopy(document.tables[0]._tbl.tblPr)
    for node in list(body):
        if node.tag != qn('w:sectPr'):
            body.remove(node)

    def paragraph(style, runs=None):
        p = document.add_paragraph(style=style)
        if runs:
            write_runs(p, runs)
        return p

    p = paragraph('Normal', [dict(text='Original scientific paper', bold=True)])
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.keep_with_next = True
    mark = p.add_run()
    mark.font.superscript = True
    ref = OxmlElement('w:footnoteReference')
    ref.set(qn('w:id'), '1')
    ref.set(qn('w:customMarkFollows'), '1')
    mark._r.append(ref)
    mark.add_text('*')
    paragraph('Heading 1', [dict(text=data['title'])])
    paragraph('Heading 2', [dict(text=data['author'])])
    paragraph('Affiliation', [dict(text=data['institution'])])
    p = paragraph('ORCID', [dict(text='ORCID iDs:\t' + data['author'] + '\t')])
    hyperlink(p, 'https://orcid.org/' + data['orcid'], 'https://orcid.org/' + data['orcid'])
    paragraph('Abstract', [dict(text='Abstract. ', bold=True, italic=False)] + data['abstract'])
    paragraph('Key words', [dict(text='Key words: ', bold=True, italic=False)] + data['keywords'])
    replace_footnote(document, data)

    section = subsection = subsubsection = figure = table_number = equation = 0
    for block in data['blocks']:
        kind = block['kind']
        if kind == 'section':
            section += 1
            subsection = subsubsection = 0
            paragraph('Heading 3', [dict(text=f'{section}. ')] + block['runs'])
        elif kind in ('subsection', 'subsubsection'):
            if kind == 'subsection':
                subsection += 1
                subsubsection = 0
                label, style = f'{section}.{subsection}. ', 'Heading 4'
            else:
                subsubsection += 1
                label, style = f'{section}.{subsection}.{subsubsection}. ', 'Heading 5'
            paragraph(style, [dict(text=label)] + block['runs'])
        elif kind == 'paragraph':
            paragraph('Normal', block['runs'])
        elif kind == 'equation':
            equation += 1
            p = paragraph('Equation')
            p.add_run('\t')
            p._p.append(math(block['math']))
            p.add_run(f'\t({equation})')
        elif kind == 'figure':
            figure += 1
            p = paragraph('Figure')
            p.add_run().add_picture(str(figure_assets(block['image'])), width=Cm(WIDTH_CM))
            paragraph('Figure title', [dict(text=f'Fig. {figure} ')] + block['caption'])
        elif kind == 'table':
            table_number += 1
            paragraph('Table title', [dict(text=f'Table {table_number} ', bold=True)] + block['caption'])
            ncols = len(block['columns'])
            table = document.add_table(rows=len(block['rows']), cols=ncols)
            props = deepcopy(table_properties)
            props.find(qn('w:tblW')).set(qn('w:w'), str(round(Cm(WIDTH_CM).twips)))
            props.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
            table._tbl.replace(table._tbl.tblPr, props)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            weights = [5.0] * ncols
            for row in block['rows']:
                col = 0
                for cell in row:
                    length = len(''.join(r.get('text', r.get('math', '')) for r in cell['runs']))
                    if cell['span'] == 1:
                        weights[col] = max(weights[col], min(length, 23) + 2)
                    col += cell['span']
            widths = [Cm(WIDTH_CM * w / sum(weights)) for w in weights]
            for column, width in zip(table.columns, widths):
                column.width = width
            for ri, row in enumerate(block['rows']):
                target_row = table.rows[ri]
                trpr = target_row._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if ri < block['header_rows']:
                    trpr.append(OxmlElement('w:tblHeader'))
                for cell, width in zip(target_row.cells, widths):
                    cell.width = width
                col = 0
                for cell in row:
                    target = target_row.cells[col]
                    if cell['span'] > 1:
                        target = target.merge(target_row.cells[col + cell['span'] - 1])
                    target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = target.paragraphs[0]
                    p.style = document.styles['Normal']
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.keep_with_next = ri < len(block['rows']) - 1
                    p.paragraph_format.keep_together = True
                    p.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT,
                                   'c': WD_ALIGN_PARAGRAPH.CENTER}[block['columns'][col]]
                    write_runs(p, cell['runs'])
                    for run in p.runs:
                        run.font.size = Pt(9)
                    if ri == block['header_rows'] - 1:
                        borders = OxmlElement('w:tcBorders')
                        border = OxmlElement('w:bottom')
                        border.set(qn('w:val'), 'single')
                        border.set(qn('w:sz'), '4')
                        borders.append(border)
                        target._tc.get_or_add_tcPr().append(borders)
                    col += cell['span']
            if block.get('notes'):
                p = paragraph('Normal', block['notes'])
                p.paragraph_format.first_line_indent = Cm(0)
                for run in p.runs:
                    run.font.size = Pt(8)
        elif kind == 'list_item':
            number = block['number']
            label = (f'{number}. ' if block.get('ordered', True) else '\u2022 ') if number is not None else ''
            paragraph('Normal', [dict(text=label)] + block['runs'])
        else:
            raise ValueError('Unsupported Facta block: ' + kind)
    paragraph('Heading 3', [dict(text='References')])
    database = shared.bibtexparser.loads((ROOT / 'references.bib').read_text(encoding='utf-8'))
    entries = database.entries_dict
    for index, reference in enumerate(data['references'], 1):
        paragraph('References', [dict(text=f'[{index}]\t')] + reference_runs(reference, entries[reference['key']]))
    document.core_properties.title = data['title']
    document.core_properties.author = data['author']
    document.core_properties.subject = 'Facta Universitatis, Series: Electronics and Energetics'
    document.core_properties.comments = 'Generated from shared LaTeX/BibTeX sources using Facta_Template.docx.'
    # Remove unused template equation objects and illustration relationships.
    used = {v for el in body.iter() for k, v in el.attrib.items()
            if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
    for rid, rel in list(document.part.rels.items()):
        if rid not in used and rel.reltype.rsplit('/', 1)[-1] in ('image', 'oleObject'):
            document.part.drop_rel(rid)
    document.save(temporary)


def validate(path, data):
    d = Document(path)
    with ZipFile(TEMPLATE) as original, ZipFile(path) as output:
        source = parse_xml(original.read('word/document.xml'))
        result = parse_xml(output.read('word/document.xml'))
        for tag in ('pgSz', 'pgMar', 'cols'):
            assert source.find('.//' + qn('w:' + tag)).attrib == result.find('.//' + qn('w:' + tag)).attrib
        assert original.read('word/styles.xml') == output.read('word/styles.xml')
        for name in original.namelist():
            if re.match(r'word/(?:header|footer)\d+\.xml$', name):
                assert original.read(name) == output.read(name), 'Header/footer template changed'
        assert not result.findall('.//' + qn('w:object')), 'Legacy equation objects remain'
        actual_math = result.findall('.//' + qn('m:oMath'))
        expected_math = list(shared.source_math(data))
        assert len(actual_math) == len(expected_math)
        def structure(node):
            return node.tag, sorted(node.attrib.items()), node.text or '', [structure(c) for c in node]
        for actual, tex in zip(actual_math, expected_math):
            assert structure(actual) == structure(math(tex)), 'Native equation structure mismatch'
        assert len(result.findall('.//' + qn('wp:inline'))) == sum(b['kind'] == 'figure' for b in data['blocks'])
        media = [n for n in output.namelist() if n.startswith('word/media/')]
        assert len(media) == 3 and all(n.endswith('.jpg') for n in media), media
        assert data['institution'] in ''.join(result.itertext())
        assert 'https://orcid.org/' + data['orcid'] in ''.join(result.itertext())
        footnotes = output.read('word/footnotes.xml').decode()
        assert data['email'] in footnotes and data['author'] in footnotes
    assert len(d.tables) == sum(b['kind'] == 'table' for b in data['blocks'])
    assert len([p for p in d.paragraphs if p.style.style_id == 'References']) == len(data['references'])
    assert d.styles['Normal'].font.size.pt == 10
    assert d.styles['Heading 1'].font.size.pt == 12
    assert d.styles['Heading 2'].font.size.pt == 10
    assert d.styles['Abstract'].font.size.pt == 9
    assert d.styles['Normal'].font.name == 'Times New Roman'
    for p in d.paragraphs:
        if p.style.style_id == 'Equation':
            assert p.text.startswith('\t') and re.search(r'\t\(\d+\)$', p.text)
    tabs = d.styles['Equation'].paragraph_format.tab_stops
    assert [round(t.position.cm, 2) for t in tabs] == [6.25, 12.7]


def main():
    shared.ROOT = ROOT
    data = shared.read_manuscript()
    details = json.loads((ROOT / 'content/author_details.json').read_text(encoding='utf-8'))
    data['orcid'] = details['orcid']
    assert re.fullmatch(r'\d{4}-\d{4}-\d{4}-\d{3}[\dX]', data['orcid'])
    total = 0
    for digit in data['orcid'].replace('-', '')[:-1]:
        total = (total + int(digit)) * 2
    check = (12 - total % 11) % 11
    assert data['orcid'][-1] == ('X' if check == 10 else str(check)), 'Invalid ORCID checksum'
    data['institution'], separator, data['email'] = data['affiliation'].partition('; E-mail: ')
    assert separator and '@' in data['email']
    keywords = ''.join(r.get('text', '') for r in data['keywords']).rstrip('.')
    assert len(keywords.split(';')) <= 6, 'Facta permits at most six keywords'
    data['source_sha256']['content/author_details.json'] = fingerprint(ROOT / 'content/author_details.json')
    template_hash = fingerprint(TEMPLATE)
    BUILD.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='facta-', dir=BUILD) as directory:
        temporary = Path(directory) / OUTPUT.name
        render(data, temporary)
        validate(temporary, data)
        assert fingerprint(TEMPLATE) == template_hash
        for name, digest in data['source_sha256'].items():
            assert fingerprint(ROOT / name) == digest, 'Source changed during build: ' + name
        temporary.replace(OUTPUT)
    report = dict(output=str(OUTPUT), generator='Official Facta template; shared LaTeX/BibTeX sources',
        template_sha256=template_hash, source_sha256=data['source_sha256'],
        tables=sum(b['kind'] == 'table' for b in data['blocks']),
        figures=sum(b['kind'] == 'figure' for b in data['blocks']),
        equations=sum(b['kind'] == 'equation' for b in data['blocks']),
        omml_expressions=len(list(shared.source_math(data))), references=len(data['references']),
        orcid=data['orcid'], figure_format='300 dpi quality-100 JPEG; companion TIFF files', warnings=[],
        manual_review=['Fill publication header/footer placeholders and received date.',
                       'Native Word equations use automatic layout; the legacy Equation Editor screenshot is not an OLE requirement.'])
    (BUILD / 'build_report.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    (BUILD / 'manuscript.json').write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print('Created and validated', OUTPUT, 'using the official Facta template.')


if __name__ == '__main__':
    main()
