"""Normalize the converted official template to its written formatting rules."""
from hashlib import sha256
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent


def main():
    document = Document(ROOT / 'build/template_review/converted_template.docx')
    normal = document.styles['Normal']

    def style(name, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
              before=0, after=0, indent=0, bold=False, italic=False, keep=False):
        s = document.styles[name]
        if name != 'Normal':
            s.base_style = normal
        s.font.name = 'Times New Roman'
        s.font.size = Pt(size)
        s.font.bold, s.font.italic = bold, italic
        s.font.all_caps = False
        s.font.small_caps = False
        p = s.paragraph_format
        p.alignment = align
        p.space_before, p.space_after = Pt(before), Pt(after)
        p.first_line_indent = Cm(indent)
        p.left_indent = p.right_indent = Cm(0)
        p.line_spacing = 1.0
        p.keep_with_next = keep
        p.keep_together = keep
        p.widow_control = True
        return s

    center, left = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    style('Normal', indent=.5)
    style('Heading 1', 12, center, 30, 0, bold=True, keep=True).font.all_caps = True
    style('Heading 2', 10, center, 0, 12, bold=True, keep=True)
    style('Affiliation', 10, center, 0, 18, keep=True)
    orcid = style('ORCID', 8, left, keep=True)
    orcid.paragraph_format.left_indent = Cm(.75)
    orcid.paragraph_format.tab_stops.clear_all()
    orcid.paragraph_format.tab_stops.add_tab_stop(Cm(2.4))
    orcid.paragraph_format.tab_stops.add_tab_stop(Cm(6.2))
    abstract = style('Abstract', 9, before=12, italic=True)
    # Actual supplied style uses 425 twips = 0.75 cm, not the prose typo "0.75 pt".
    abstract.paragraph_format.left_indent = abstract.paragraph_format.right_indent = Cm(.75)
    keywords = style('Key words', 9, before=6, italic=True)
    keywords.paragraph_format.left_indent = keywords.paragraph_format.right_indent = Cm(.75)
    style('Heading 3', 10, center, 24, 6, keep=True).font.small_caps = True
    for name, bold, italic in [('Heading 4', True, False), ('Heading 5', False, True)]:
        s = style(name, 10, left, 12, 6, bold=bold, italic=italic, keep=True)
        s.paragraph_format.left_indent = Pt(5)
    eq = style('Equation', 10, left, 6, 6)
    eq.paragraph_format.tab_stops.clear_all()
    eq.paragraph_format.tab_stops.add_tab_stop(Cm(6.25), WD_TAB_ALIGNMENT.CENTER)
    eq.paragraph_format.tab_stops.add_tab_stop(Cm(12.7), WD_TAB_ALIGNMENT.RIGHT)
    style('Figure', 10, center, 12, 0, keep=True)
    style('Figure title', 10, center, 6, 12, bold=True).paragraph_format.keep_together = True
    style('Table title', 10, center, 12, 6, keep=True)
    refs = style('References', 8, left)
    refs.paragraph_format.left_indent = Cm(.5)
    refs.paragraph_format.first_line_indent = Cm(-.5)
    refs.paragraph_format.tab_stops.add_tab_stop(Cm(.5))
    # Explicit numeric labels use the resolved BibTeX citation order.
    num = refs.element.find('.//' + qn('w:numPr'))
    if num is not None:
        num.getparent().remove(num)
    style('Acknowledgement', 9, before=12, italic=True)
    style('footnote text', 8, left)
    defaults = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    for tag, value in [('sz', '20'), ('szCs', '20')]:
        node = defaults.find(qn('w:' + tag))
        if node is None:
            node = OxmlElement('w:' + tag)
            defaults.append(node)
        node.set(qn('w:val'), value)
    language = defaults.find(qn('w:lang'))
    language.set(qn('w:val'), 'en-US')
    # Replace publication metadata with clearly editable placeholders, retaining
    # the first/odd/even header and footer structure and page fields.
    for section in document.sections:
        for header, text in [(section.even_page_header, '[Running author(s)]'),
                             (section.header, '[Running title]')]:
            p = header.paragraphs[0]
            p.clear()
            p.add_run(text + '\t')
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE')
            p._p.append(field)
        first = section.first_page_header
        for p in first.paragraphs:
            p.clear()
        first.paragraphs[0].add_run('FACTA UNIVERSITATIS').bold = True
        first.paragraphs[1].add_run('Series: Electronics and Energetics Vol. [volume], No. [issue], [month year], pp. [pages]')
        first.paragraphs[2].add_run('https://doi.org/[to be assigned]')
        for p in first.paragraphs:
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.all_caps = False
        footer = section.first_page_footer.paragraphs[0]
        footer.clear()
        footer.add_run('\u00a9 [year] by University of Ni\u0161, Serbia | Creative Commons License: CC BY-NC-ND')
        for run in footer.runs:
            run.font.size = Pt(7)
            run.font.all_caps = False
    output = ROOT / 'Facta_Template.docx'
    document.save(output)
    manifest = dict(source='fu_elecenerg_instructions_for_authors.doc',
        source_sha256=sha256((ROOT / 'fu_elecenerg_instructions_for_authors.doc').read_bytes()).hexdigest(),
        template=output.name, template_sha256=sha256(output.read_bytes()).hexdigest(),
        policy='Written instructions override conflicting stored styles; page geometry retained.',
        differences=['Title spacing normalized to 30/0 pt instead of stored 18/18.',
                     'Author normalized to 10 pt instead of stored 12.',
                     'Affiliation spacing after normalized to 18 pt instead of stored 12.',
                     'Subheading left indent normalized to 5 pt instead of about 14.',
                     'Abstract indentation uses actual template 0.75 cm (425 twips), correcting prose unit typo.',
                     'Publication header/footer metadata replaced with editable placeholders.'])
    (ROOT / 'facta_template_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    print('Prepared', output)


if __name__ == '__main__':
    main()

