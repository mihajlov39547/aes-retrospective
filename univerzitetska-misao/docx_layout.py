"""Adapt SJEE styles to uputstvo_za_autore.pdf, pages 1-3 and 7."""
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


def style(document, name, size, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
          bold=False, italic=False, before=0, after=0, indent=0, keep=False):
    styles = document.styles
    s = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = styles['Normal'] if name != 'Normal' else None
    s.font.name = 'Times New Roman'
    s.font.size = Pt(size)
    s.font.bold, s.font.italic = bold, italic
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.all_caps = False
    pf = s.paragraph_format
    pf.alignment = alignment
    pf.line_spacing = 1
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.first_line_indent = Cm(indent)
    pf.left_indent = pf.right_indent = Cm(0)
    pf.keep_with_next = keep
    pf.keep_together = False
    pf.page_break_before = False
    pf.widow_control = True
    pf.tab_stops.clear_all()
    # JReference in the SJEE template carries its own [n] numbering. We write
    # the journal's 1., 2., ... explicitly, so inherited numbering must go.
    for numbering in s.element.xpath('./w:pPr/w:numPr'):
        numbering.getparent().remove(numbering)
    return s


def apply(p, name):
    # Pandoc and the source template can contain direct formatting that wins
    # over a named style. Clear it while retaining numbering/bookmarks/math.
    pr = p._p.pPr
    if pr is not None:
        for child in list(pr):
            if child.tag != qn('w:numPr'):
                pr.remove(child)
    p.style = name
    for r in p._p.xpath('.//w:r'):
        props = r.find(qn('w:rPr'))
        if props is not None:
            for child in list(props):
                if child.tag not in (qn('w:b'), qn('w:i'), qn('w:vertAlign')):
                    props.remove(child)
        # Also applies to runs inside hyperlinks.
        if props is None:
            props = OxmlElement('w:rPr')
            r.insert(0, props)
        fonts = OxmlElement('w:rFonts')
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), 'Times New Roman')
        props.append(fonts)


def insert_after(element, text, document, name):
    new = OxmlElement('w:p')
    element.addnext(new)
    p = Paragraph(new, document._body)
    p.add_run(text)
    apply(p, name)
    return p


def format_docx(path, kind):
    d = Document(path)
    d.core_properties.author = ''
    d.core_properties.last_modified_by = ''
    d.core_properties.title = d.paragraphs[0].text.strip()
    left, center, justify = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
    style(d, 'Normal', 11)
    style(d, 'JText', 11, indent=.5)
    style(d, 'JTitle', 13, center, bold=True, after=10, keep=True)
    style(d, 'JAbstract', 10)
    style(d, 'UMAbstractHeading', 10, left, bold=True, before=4, after=3, keep=True)
    style(d, 'UMKeywords', 10, left, bold=True, before=5, after=12)
    style(d, 'JSubtitle', 12, justify, bold=True, before=10, after=6, keep=True)
    style(d, 'JSubSubTitle', 11, left, bold=True, before=8, after=4, keep=True)
    style(d, 'JReference', 10, justify, after=4)
    style(d, 'JTableText', 9, left)
    style(d, 'UMTableNumber', 10, left, bold=True, before=8, after=2, keep=True)
    style(d, 'UMTableTitle', 10, left, italic=True, after=4, keep=True)
    style(d, 'UMTableSource', 10, left, before=3, after=8)
    style(d, 'JEquation', 11, center, before=6, after=6)
    style(d, 'UMFigureCaption', 10, left, italic=True, after=3, keep=True)
    style(d, 'UMFigureNumber', 10, left, bold=True, after=2, keep=True)
    style(d, 'JAuthor', 10, center, bold=True, after=8)
    style(d, 'UMCoverDetail', 10, center, after=8)
    style(d, 'UMCoverContact', 10, center, italic=True, after=8)
    for section in d.sections:
        section.page_width, section.page_height = Cm(17.6), Cm(25.1)
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)
        section.gutter = Cm(0)
        for tag in ('headerReference', 'footerReference', 'pgNumType'):
            for el in section._sectPr.findall(qn('w:' + tag)):
                section._sectPr.remove(el)
        cols = section._sectPr.find(qn('w:cols'))
        if cols is not None:
            cols.clear()
            cols.set(qn('w:num'), '1')
    if kind == 'cover':
        for i, p in enumerate(d.paragraphs):
            name = 'JTitle' if i == 0 else 'JAuthor' if i == 1 else (
                'UMCoverContact' if p.text.strip().startswith(('E-mail:', 'ORCID:')) else 'UMCoverDetail')
            apply(p, name)
            if i == 0:
                p.paragraph_format.space_before = Cm(2.5)
                p.paragraph_format.space_after = Cm(1.8)
            if re.fullmatch(r'\s*\d{4}\s*', p.text):
                p.paragraph_format.space_before = Cm(4)
        d.save(path)
        return {}

    abstract = False
    references = False
    reference_number = 0
    abstract_words = {}
    equation_number = 0
    figure_number = 0
    paragraphs = list(d.paragraphs)
    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        if i == 0 or text.startswith(('A VALIDATED METHODOLOGICAL FRAMEWORK', 'CPU/GPU AES ENCRYPTION')):
            apply(p, 'JTitle')
        elif text in ('Apstrakt', 'Abstract'):
            abstract = text
            apply(p, 'UMAbstractHeading')
        elif text.startswith(('Ključne reči:', 'Keywords:')):
            abstract = False
            apply(p, 'UMKeywords')
        elif abstract:
            apply(p, 'JAbstract')
            abstract_words[abstract] = len(text.split())
        elif text in ('Literatura', 'SUMMARY') or re.match(r'^\d+\t', text):
            references = text == 'Literatura'
            apply(p, 'JSubtitle')
            for node in p._p.xpath('.//w:t'):
                node.text = node.text.upper() if node.text else node.text
            # Replace Pandoc's tab separator with punctuation and a space.
            for tab in p._p.xpath('.//w:tab'):
                t = OxmlElement('w:t')
                t.set(qn('xml:space'), 'preserve')
                t.text = '. '
                tab.getparent().replace(tab, t)
            if text == 'SUMMARY':
                p.paragraph_format.alignment = center
                p.paragraph_format.page_break_before = True
        elif re.match(r'^\d+\.\d+\t', text):
            apply(p, 'JSubSubTitle')
            for tab in p._p.xpath('.//w:tab'):
                t = OxmlElement('w:t')
                t.set(qn('xml:space'), 'preserve')
                t.text = '. '
                tab.getparent().replace(tab, t)
        elif references:
            reference_number += 1
            apply(p, 'JReference')
            for num in p._p.xpath('./w:pPr/w:numPr'):
                num.getparent().remove(num)
            r = p.add_run(f'{reference_number}. ')
            p._p.insert(1, r._r)
            p.paragraph_format.left_indent = Cm(.55)
            p.paragraph_format.first_line_indent = Cm(-.55)
        elif p._p.xpath('.//m:oMathPara'):
            apply(p, 'JEquation')
            equation_number += 1
            p.add_run(f' ({equation_number})')
        elif p._p.xpath('.//w:drawing'):
            apply(p, 'JText')
            p.paragraph_format.alignment = center
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
        elif i > 0 and paragraphs[i - 1]._p.xpath('.//w:drawing'):
            figure_number += 1
            apply(p, 'UMFigureCaption')
            p.insert_paragraph_before(f'Slika {figure_number}', 'UMFigureNumber')
            insert_after(p._p, 'Izvor: Eksperimentalni rezultati autora.', d, 'UMTableSource')
        else:
            apply(p, 'JText')

    for shape in d.inline_shapes:
        if shape.width > Cm(12.6):
            shape.height = int(shape.height * Cm(12.6) / shape.width)
            shape.width = Cm(12.6)

    # Explicit widths fit the 12.6 cm printable width, including merged cells.
    widths = {6: [1.5, 2.3, 2.2, 2.2, 2.2, 2.2], 4: [2.0, 2.4, 4.1, 4.1]}
    for number, table in enumerate(d.tables, 1):
        caption = Paragraph(table._tbl.getprevious(), d._body)
        if not caption.text.strip():
            raise ValueError(f'Table {number} is missing its caption')
        caption.insert_paragraph_before(f'Tabela {number}', 'UMTableNumber')
        apply(caption, 'UMTableTitle')
        source = ('Izvor: Eksperimentalni rezultati autora.' if number == 1 else
                  'Izvor: Proračun autora na osnovu eksperimentalnih rezultata; normalizovani troškovi.')
        insert_after(table._tbl, source, d, 'UMTableSource')
        table.autofit = False
        grid = widths.get(len(table.columns), [12.6 / len(table.columns)] * len(table.columns))
        for col, width in zip(table.columns, grid):
            col.width = Cm(width)
        props = table._tbl.tblPr
        for tag in ('tblInd', 'tblBorders', 'tblCellMar', 'tblW'):
            for el in props.findall(qn('w:' + tag)):
                props.remove(el)
        total = OxmlElement('w:tblW')
        total.set(qn('w:w'), str(round(Cm(12.6).twips)))
        total.set(qn('w:type'), 'dxa')
        props.append(total)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'bottom', 'insideH', 'left', 'right', 'insideV'):
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single' if edge in ('top', 'bottom', 'insideH') else 'nil')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:color'), '000000')
            borders.append(el)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge, value in [('top', 35), ('bottom', 35), ('left', 35), ('right', 35)]:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:w'), str(value))
            el.set(qn('w:type'), 'dxa')
            margins.append(el)
        props.append(margins)
        for row_idx, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if row_idx == 0 or (number == 3 and row_idx == 1):
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            offset = 0
            for tc in row._tr.tc_lst:
                span = tc.grid_span
                tc.width = Cm(sum(grid[offset:offset + span]))
                offset += span
            seen = set()
            for col_idx, cell in enumerate(row.cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for shading in cell._tc.xpath('./w:tcPr/w:shd'):
                    shading.getparent().remove(shading)
                for p in cell.paragraphs:
                    apply(p, 'JTableText')
                    p.paragraph_format.alignment = left if col_idx == 0 else center
                    p.paragraph_format.keep_with_next = row_idx < len(table.rows) - 1
                    if row_idx == 0 or (number == 3 and row_idx == 1):
                        for r in p.runs:
                            r.bold = True
    for label, count in abstract_words.items():
        if not 100 <= count <= 200:
            raise ValueError(f'{label} has {count} words; the journal requires 100-200')
    d.save(path)
    return {'abstract_words': abstract_words, 'page_cm': [17.6, 25.1], 'margins_cm': 2.5}
