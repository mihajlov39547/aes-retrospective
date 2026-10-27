r"""Regression checks for manuscript constructs added in October 2026.

Run: .\.venv-docx\Scripts\python.exe -m unittest discover -s tests
"""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from docx.oxml.ns import qn
import fill_sjee_template_docx as g


class AdapterTests(unittest.TestCase):
    def test_bibliography_accents_and_doi(self):
        refs = {r['key']: r['text'] for r in g.bibliography()}
        entry = refs['schonberger2011gpu']
        self.assertIn('G. Sch\u00f6nberger, J. Fu\u00df', entry)
        self.assertIn('Lecture Notes in Computer Science', entry)
        self.assertIn('10.1007/978-3-642-24712-5_16', entry)
        self.assertEqual(g.plain(g.parse(r'Sch{\"o}nberger and Fu{\ss}')),
                         'Sch\u00f6nberger and Fu\u00df')

    def test_numbered_list_keeps_all_paragraphs(self):
        result = g.blocks(r'\begin{enumerate}\item First.\item \textbf{Second.}\end{enumerate}')
        self.assertEqual([b['number'] for b in result], [1, 2])
        self.assertTrue(result[1]['runs'][0]['bold'])

    def test_current_equation_symbols_and_structures(self):
        data = g.read_manuscript()
        self.assertEqual(len(list(g.source_math(data))), 61)
        equations = [g.omml(b['math']) for b in data['blocks'] if b['kind'] == 'equation']
        text = lambda node: ''.join(n.text or '' for n in node.iter(qn('m:t')))
        self.assertIn('\u2208', text(equations[0]))
        self.assertEqual(text(equations[1]).count('\u00d7'), 3)
        for unit in ('H2D', 'kernel', 'D2H', 'overhead'):
            self.assertIn(unit, text(equations[2]))
        self.assertEqual(len(list(equations[3].iter(qn('m:f')))), 1)
        self.assertEqual(len(list(equations[4].iter(qn('m:acc')))), 2)
        self.assertEqual(len(list(equations[4].iter(qn('m:d')))), 2)
        self.assertEqual(len(list(equations[5].iter(qn('m:acc')))), 3)
        self.assertIn('AES-NI', text(equations[5]))
        self.assertIn('\u2026', text(equations[5]))

    def test_unordered_list_uses_native_bullet_definition(self):
        result = g.blocks(r'\begin{itemize}\item One.\item Two.\end{itemize}')
        self.assertTrue(all(not b['ordered'] for b in result))
        doc = Document(g.SOURCE)
        first = g.list_numbering(doc, True)
        second = g.list_numbering(doc, True)
        self.assertNotEqual(first, second)
        number = next(n for n in doc.part.numbering_part.element
                      if n.get(qn('w:numId')) == str(second))
        self.assertEqual(number.find('.//' + qn('w:startOverride')).get(qn('w:val')), '1')

    def test_widetilde_is_native_accent(self):
        node = g.omml(r'\frac{\widetilde{T}_{CPU}}{\widetilde{T}_{GPU}}')
        accents = list(node.iter(qn('m:acc')))
        self.assertEqual(len(accents), 2)
        self.assertTrue(all(a.find(qn('m:accPr')).find(qn('m:chr')).get(qn('m:val')) == '\u0303' for a in accents))

    def test_forward_reference_and_citation_order(self):
        source = r'\section{Test}See \ref{tab:later}, \cite{second,first}.\begin{table}\caption{Later}\label{tab:later}\end{table}'
        data = {'references': [{'key':'first'}, {'key':'second'}, {'key':'uncited'}],
                'runs': g.inline(g.parse(r'\ref{tab:later} \cite{second,first}'))}
        g.resolve_links(data, {'content/01_test.tex': source})
        self.assertEqual(data['citation_numbers'], {'second':1, 'first':2, 'uncited':3})
        self.assertEqual(''.join(r['text'] for r in data['runs']), '1 [1, 2]')

    def test_missing_reference_fails(self):
        data = {'references': [], 'runs': g.inline(g.parse(r'\ref{missing}'))}
        with self.assertRaisesRegex(ValueError, 'Unknown cross-reference'):
            g.resolve_links(data, {})

    def test_unknown_construct_fails(self):
        with self.assertRaisesRegex(ValueError, 'Unsupported'):
            g.blocks(r'Important \unknown{content}')

    def test_tabularx_and_table_note(self):
        source = r'''\begin{table}\caption{Platforms}
        \begin{tabularx}{\linewidth}{@{}l >{\raggedright\arraybackslash}X@{}}
        \toprule ID & Hardware \\ \midrule A & CPU \\ \addlinespace
        B & GPU \\ \bottomrule\end{tabularx}
        \par\smallskip\parbox{\linewidth}{\footnotesize Note.}\end{table}'''
        result = g.blocks(source)[0]
        self.assertEqual(result['columns'], 'll')
        self.assertEqual(len(result['rows']), 3)
        self.assertEqual(result['notes'][0]['text'], 'Note.')

    def test_table_reference_word_and_number_are_bold(self):
        source = r'\section{Test}\begin{table}\caption{Test}\label{tab:test}\end{table}'
        data = {'references': [], 'runs': g.inline(g.parse(r'See Table~\ref{tab:test}.'))}
        g.resolve_links(data, {'content/01_test.tex': source})
        bold = ''.join(r.get('text', '') for r in data['runs'] if r.get('bold'))
        self.assertEqual(bold, 'Table\u00a01')

    def test_native_roman_units_and_math_size(self):
        node = g.omml(r'T_{\mathrm{GPU}} + 1\,\mathrm{KiB}')
        texts = [r.text for r in node.iter(qn('m:t'))]
        self.assertIn('GPU', texts)
        self.assertIn('KiB', texts)
        self.assertTrue(all(r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val')) == '22'
                            for r in node.iter(qn('m:r'))))

    def test_current_manuscript_export(self):
        data = g.read_manuscript()
        self.assertEqual(sum(b['kind'] == 'table' for b in data['blocks']), 6)
        self.assertEqual(sum(b['kind'] == 'figure' for b in data['blocks']), 2)
        self.assertEqual(sum(b['kind'] == 'equation' for b in data['blocks']), 6)
        self.assertEqual(len(data['references']), 23)
        self.assertIn('content/focused_results_table.tex', data['source_sha256'])
        self.assertEqual(data['labels']['fig:bulk-throughput'], '2')
        with TemporaryDirectory(dir=g.BUILD) as directory:
            path = Path(directory) / 'check.docx'
            self.assertEqual(g.render(data, path), [])
            g.validate(path, data)
            doc = Document(path)
            items = [p for p in doc.paragraphs if p.style.style_id == 'JText-numberedList']
            self.assertEqual(len(items), 3)
            self.assertTrue(all(p._p.find('.//' + qn('w:numPr')) is not None for p in items))
            self.assertFalse(any(p.text.startswith('1. ') for p in items))
            self.assertTrue(doc.paragraphs[0]._p.findall('.//' + qn('w:br')))
            equations = [p for p in doc.paragraphs if p.style.style_id == 'JEquation']
            self.assertTrue(all(len(p._p.findall('.//' + qn('w:tab'))) == 2 for p in equations))
            self.assertTrue(all(p._p.findall('.//' + qn('m:oMath')) for p in equations))
            for table in doc.tables:
                for row in table.rows:
                    self.assertIsNotNone(row._tr.find('.//' + qn('w:cantSplit')))


if __name__ == '__main__':
    unittest.main()
