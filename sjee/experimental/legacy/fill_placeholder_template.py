r"""Experimental, template-specific filling; never part of the primary PDF build.

Run with .\.venv-docx\Scripts\python.exe fill_sjee_template_docx.py.
Requires python-docx; see requirements_docx_experiment.txt.
Only placeholders are supplied. No LaTeX parsing or bibliography automation.
"""

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from shutil import copyfile
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "SJEE_Template.docx"
OUTPUT = ROOT / "main_sjee_filled.docx"
# Index mapping below is intentionally bound to the inspected template version.
EXPECTED_SHA256 = "51ab1782847fabe6910f7fa4af264bc25c3bad8585f0db54d31fe963941e578d"
TITLE = (
    "Validation-First CPU/GPU Microbenchmarking of AES Implementations "
    "in the Context of Symmetric Cryptography Evolution"
)
AFFILIATION = "Singidunum University, Belgrade, Serbia; ORCID: 0009-0006-7979-0341"


def fill_paragraph(node, owner, text, *, bold_prefix=None):
    """Retain the original pPr/style; replace text, not the style definition."""
    paragraph = Paragraph(node, owner)
    # Preserve a representative non-footnote run's direct formatting.
    original = next((run for run in paragraph.runs if run.text.strip()), None)
    rpr = deepcopy(original._r.rPr) if original is not None and original._r.rPr is not None else None
    paragraph.clear()  # Retains paragraph properties; removes sample objects too.
    if bold_prefix is not None:
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text).bold = False
    else:
        run = paragraph.add_run(text)
        if rpr is not None:
            run._r.insert(0, rpr)
    return paragraph


def fill_footnote(document):
    """python-docx has no public footnote editor; make one bounded XML edit."""
    part = next(p for p in document.part.package.parts if str(p.partname) == "/word/footnotes.xml")
    root = parse_xml(part.blob)
    for note in list(root):
        note_id = note.get(qn("w:id"))
        if note_id == "2":
            root.remove(note)  # Second sample author is no longer present.
        elif note_id == "1":
            node = note.find(qn("w:p"))
            marker = deepcopy(node.find(".//" + qn("w:footnoteRef")).getparent())
            for child in list(node):
                if child.tag != qn("w:pPr"):
                    node.remove(child)
            node.append(marker)
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.set(qn("xml:space"), "preserve")
            text.text = " " + AFFILIATION
            run.append(text)
            node.append(run)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def fill_table(table):
    """Preserve the sample's five-column grid/merges/borders; only fill text."""
    replacements = [
        ["", "Placeholder A", "Placeholder B"],
        ["", "Measurements", "Measurements"],
        ["Case", "Time [s]", "Size [byte]", "Time [s]", "Size [byte]"],
        ["A", "—", "—", "—", "—"],
        ["B", "—", "—", "—", "—"],
    ]
    # Iterate physical cells, not row.cells, which repeats merged cells.
    for row, values in zip(table._tbl.tr_lst, replacements, strict=True):
        for cell, value in zip(row.tc_lst, values, strict=True):
            paragraphs = cell.findall(qn("w:p"))
            fill_paragraph(paragraphs[0], table, value)
            for extra in paragraphs[1:]:
                cell.remove(extra)


def verify_preservation(source, output):
    """Fail if style definitions, page setup or headers/footers were changed."""
    with ZipFile(source) as original, ZipFile(output) as filled:
        names = ["word/styles.xml", "word/numbering.xml"]
        names += [n for n in original.namelist() if n.startswith(("word/header", "word/footer")) and n.endswith(".xml")]
        for name in names:
            before = etree.tostring(parse_xml(original.read(name)), method="c14n")
            after = etree.tostring(parse_xml(filled.read(name)), method="c14n")
            if before != after:
                raise RuntimeError(f"Unexpected change to preserved part: {name}")
        def section(archive):
            root = parse_xml(archive.read("word/document.xml"))
            return etree.tostring(root.find(".//" + qn("w:sectPr")), method="c14n")
        if section(original) != section(filled):
            raise RuntimeError("Unexpected change to section/page settings")


def main():
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    if source_hash != EXPECTED_SHA256:
        raise RuntimeError("The source template changed. Reinspect its structure before adapting this experimental script.")
    if not (ROOT / "figures/placeholder.png").is_file():
        raise FileNotFoundError("Missing figures/placeholder.png")
    if OUTPUT.is_symlink() or (OUTPUT.exists() and SOURCE.samefile(OUTPUT)):
        raise RuntimeError("Output must be a separate file, not a link to the source.")
    copyfile(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    document.core_properties.title = TITLE
    document.core_properties.author = "Marko Mihajlovic"
    document.core_properties.subject = "Experimental SJEE placeholder manuscript"
    body = document._element.body
    nodes = list(body)
    # These are zero-based BODY child indices, including the table/bookmark.
    expected = {0: "JTitle", 1: "JAuthor", 2: "JAbstract", 3: "JAbstract",
                4: "JSubtitle", 10: "JSubSubTitle", 13: "JEquation",
                25: "JFigure", 42: "JFigure", 52: "JReference"}
    for index, style in expected.items():
        if Paragraph(nodes[index], document._body).style.style_id != style:
            raise RuntimeError(f"Template structure mismatch at body child {index}")
    author_marker = deepcopy(nodes[1].find(".//" + qn("w:footnoteReference")).getparent())

    replacements = {
        0: TITLE,
        1: "Marko Mihajlovic",
        4: "1\tIntroduction",
        5: "Placeholder introduction. Describe the motivation, scope, and organization of the manuscript. The citation here demonstrates the numeric reference format [1].",
        8: "2\tRelated Work and Research Gap",
        9: "Placeholder related work. Summarize relevant literature and state the research gap after the manuscript content and sources have been prepared.",
        17: "3\tMethodology",
        18: "Placeholder methodology. Describe the experimental procedure and define the quantities used in the analysis.",
        10: "3.1\tExample equation",
        12: "Equation (1) is reserved for a typesetting example; insert a proper Word or MathType equation manually.",
        13: "\t[Equation placeholder]\t(1)",
        16: "This is placeholder text, not an equation. Native equation editing, numbering and cross-references require manual Word/MathType cleanup.",
        31: "4\tResults",
        32: "Placeholder results. Fig. 1 demonstrates a single figure placed near its first mention.",
        40: "5\tDiscussion",
        46: "Placeholder discussion. Interpret the eventual results and describe their limitations and implications.",
        45: "6\tConclusion",
        47: "Placeholder conclusion. Summarize the eventual findings and possible future work without introducing new results.",
        48: "7\tAcknowledgments",
        49: "Placeholder acknowledgments. Replace with appropriate funding or support information, or remove the paragraph if no acknowledgments are needed.",
        51: "References",
        # JReference inherits the existing [%1] list: do not type another [1].
        52: "Placeholder Author: Placeholder reference title, 2026. Fictional entry for template testing only; replace before submission.",
    }
    for index, text in replacements.items():
        fill_paragraph(nodes[index], document._body, text)
    nodes[1].append(author_marker)
    fill_paragraph(nodes[2], document._body,
                   "Placeholder abstract. Replace this paragraph with a concise description of the purpose, methods, principal results, and contribution of the completed manuscript.",
                   bold_prefix="Abstract: ")
    fill_paragraph(nodes[3], document._body,
                   "Placeholder keyword, Second keyword, Third keyword.", bold_prefix="Keywords: ")
    image = fill_paragraph(nodes[25], document._body, "")
    image.add_run().add_picture(str(ROOT / "figures/placeholder.png"), width=Cm(10))
    fill_paragraph(nodes[29], document._body,
                   "Placeholder figure; replace with the final illustration.", bold_prefix="Fig. 1 − ")
    fill_paragraph(nodes[41], document._body,
                   " demonstrates the template's table geometry with placeholder entries and units in square brackets. No experimental data are shown.", bold_prefix="Table 1")
    caption = fill_paragraph(nodes[42], document._body,
                            "\nPlaceholder measurements; replace with verified data.", bold_prefix="Table 1")
    caption.runs[-1].italic = True
    fill_table(document.tables[0])
    fill_footnote(document)

    # Reuse/reorder original paragraph and table nodes, removing instructional
    # samples, extra figures and legacy equation objects from the visible body.
    order = [0, 1, 2, 3, 4, 5, 8, 9, 17, 18, 10, 12, 13, 16,
             31, 32, 25, 29, 41, 42, 43, 40, 46, 45, 47, 48, 49, 51, 52]
    section = nodes[-1]
    if section.tag != qn("w:sectPr"):
        raise RuntimeError("Expected a final section-properties element")
    for node in list(body):
        body.remove(node)
    for index in order:
        body.append(nodes[index])
    body.append(section)
    document.save(OUTPUT)
    verify_preservation(SOURCE, OUTPUT)
    if sha256(SOURCE.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Original template hash changed unexpectedly")
    print(f"Created experimental copy: {OUTPUT}")
    print("Verified: original unchanged; styles, numbering, headers/footers and page settings preserved.")
    print("Manual work: equations, sample production headers, final captions/references and pagination.")
    print("Primary submission workflow remains main.tex -> main.pdf.")


if __name__ == "__main__":
    main()
