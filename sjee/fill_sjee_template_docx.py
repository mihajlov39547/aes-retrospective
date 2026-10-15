"""Primary DOCX generator: copy SJEE styles, read shared sources, never use Pandoc.

The parser deliberately supports a bounded LaTeX subset and rejects unknown
constructs. build/docx/manuscript.json is a generated intermediate, not a source.
"""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from shutil import copyfile
from shutil import which
import subprocess
import textwrap
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import json
import re

import bibtexparser
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph
from latex2mathml.converter import convert as latex_mathml
from lxml import etree
from mathml2omml import convert as mathml_omml
from pylatexenc.latexwalker import (
    LatexWalker, get_default_latex_context_db, LatexCharsNode, LatexCommentNode,
    LatexMacroNode, LatexGroupNode, LatexMathNode, LatexSpecialsNode,
    LatexEnvironmentNode,
)
from pylatexenc.macrospec import MacroSpec, EnvironmentSpec
from pylatexenc.latex2text import LatexNodes2Text

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'SJEE_Template.docx'
OUTPUT = ROOT / 'main_sjee.docx'
BUILD = ROOT / 'build' / 'docx'
TEMPLATE_HASH = '51ab1782847fabe6910f7fa4af264bc25c3bad8585f0db54d31fe963941e578d'
CONTEXT = get_default_latex_context_db()
CONTEXT.add_context_category('sjee', macros=[
    MacroSpec('thanks', '{'), MacroSpec('caption', '[{'),
    MacroSpec('multicolumn', '{{{'), MacroSpec('setlength', '{{'),
    MacroSpec('url', '{'), MacroSpec('href', '{{'),
    MacroSpec('parbox', '{{'), MacroSpec('includegraphics', '[{'),
], environments=[EnvironmentSpec('tabularx', '{{')], prepend=True)


def parse(text):
    return LatexWalker(text, latex_context=CONTEXT, tolerant_parsing=False).get_latex_nodes()[0]


def argument(node, index=-1):
    return node.nodeargd.argnlist[index].nodelist


def inline(nodes, **formatting):
    result = []
    for node in nodes:
        if isinstance(node, LatexCommentNode):
            continue
        if isinstance(node, LatexCharsNode):
            result.append(dict(text=re.sub(r'\s+', ' ', node.chars), **formatting))
        elif isinstance(node, LatexGroupNode):
            result.extend(inline(node.nodelist, **formatting))
        elif isinstance(node, LatexMathNode):
            raw = node.latex_verbatim()
            result.append(dict(math=raw[len(node.delimiters[0]):-len(node.delimiters[1])], **formatting))
        elif isinstance(node, LatexSpecialsNode):
            replacements = {'~': '\u00a0', '--': '\u2013', '---': '\u2014', "``": '\u201c', "''": '\u201d'}
            if node.specials_chars not in replacements:
                raise ValueError(f'Unsupported inline special: {node.latex_verbatim()}')
            result.append(dict(text=replacements[node.specials_chars], **formatting))
        elif isinstance(node, LatexMacroNode):
            name = node.macroname
            if name in ('emph', 'textit', 'textbf', 'texttt'):
                fmt = dict(formatting)
                fmt[{'emph': 'italic', 'textit': 'italic', 'textbf': 'bold', 'texttt': 'code'}[name]] = True
                result.extend(inline(argument(node), **fmt))
            elif name in ('url', 'href'):
                result.extend(inline(argument(node), **formatting))
            elif name in ('&', '_', '%', '#', '$', '{', '}', ',', ' ', '\\'):
                result.append(dict(text={',': '\u2009', ' ': ' ', '\\': '\n'}.get(name, name), **formatting))
            elif name in ('label', 'allowbreak', 'footnotesize'):
                continue
            elif name in ('"', "'", '`', '^', '~', '=', '.', 'u', 'v', 'H', 'c',
                          'ss', 'ae', 'AE', 'oe', 'OE', 'o', 'O', 'l', 'L'):
                result.append(dict(text=LatexNodes2Text().nodelist_to_text([node]),
                                   **formatting))
            elif name in ('cite', 'ref', 'eqref'):
                result.append(dict(link_kind=name, target=plain(argument(node)), **formatting))
            else:
                raise ValueError(f'Unsupported inline command: {node.latex_verbatim()}')
        else:
            raise ValueError(f'Unsupported inline content: {node.latex_verbatim()}')
    return result


def clean(runs):
    runs = deepcopy(runs)
    while runs and 'text' in runs[0] and not runs[0]['text'].strip():
        runs.pop(0)
    while runs and 'text' in runs[-1] and not runs[-1]['text'].strip():
        runs.pop()
    if runs and 'text' in runs[0]:
        runs[0]['text'] = runs[0]['text'].lstrip()
    if runs and 'text' in runs[-1]:
        runs[-1]['text'] = runs[-1]['text'].rstrip()
    return runs


def plain(nodes):
    runs = inline(nodes)
    if any('math' in r for r in runs):
        raise ValueError('Math in metadata/BibTeX text needs explicit handling')
    return ''.join(r['text'] for r in runs).strip()


def parse_table(node):
    caption = tabular = None
    notes = []
    for child in node.nodelist:
        if isinstance(child, (LatexCommentNode, LatexCharsNode)):
            if isinstance(child, LatexCharsNode) and child.chars.strip():
                raise ValueError('Unexpected text outside tabular')
        elif isinstance(child, LatexMacroNode):
            if child.macroname == 'caption':
                caption = clean(inline(argument(child)))
            elif child.macroname == 'parbox':
                notes = clean(inline(argument(child)))
            elif child.macroname not in ('centering', 'small', 'setlength', 'label',
                                        'par', 'smallskip'):
                raise ValueError(f'Unsupported table command {child.macroname}')
        elif isinstance(child, LatexEnvironmentNode) and child.environmentname in ('tabular', 'tabularx'):
            tabular = child
        else:
            raise ValueError('Unsupported table content')
    if caption is None or tabular is None:
        raise ValueError('Each table needs a caption and a tabular')
    columns = ''.join(n.latex_verbatim() for n in argument(tabular))
    columns = re.sub(r'>\{[^}]*\}|@\{[^}]*\}|\s+', '', columns).replace('X', 'l')
    if not re.fullmatch('[lcr]+', columns):
        raise ValueError(f'Unsupported column specification: {columns}')
    raw = ''.join(n.latex_verbatim() for n in tabular.nodelist)
    raw = re.sub(r'\\cmidrule(?:\([^)]*\))?\{[^}]*\}', '', raw)
    rows, cells, pending = [], [], []
    header_rows = None
    span = 1

    def finish_cell():
        nonlocal pending, span
        cells.append({'span': span, 'runs': clean(pending)})
        pending, span = [], 1

    for child in parse(raw):
        if isinstance(child, LatexSpecialsNode) and child.specials_chars == '&':
            finish_cell()
        elif isinstance(child, LatexMacroNode) and child.macroname == '\\':
            finish_cell()
            if sum(c['span'] for c in cells) != len(columns):
                raise ValueError('Table row width mismatch')
            rows.append(cells)
            cells = []
        elif isinstance(child, LatexMacroNode) and child.macroname in ('toprule', 'midrule', 'bottomrule', 'addlinespace'):
            if child.macroname == 'midrule' and header_rows is None:
                header_rows = len(rows)
        elif isinstance(child, LatexMacroNode) and child.macroname == 'multicolumn':
            if clean(pending):
                raise ValueError('Unsupported content before multicolumn')
            span = int(plain(argument(child, 0)))
            pending = inline(argument(child, 2))
        else:
            pending.extend(inline([child]))
    if clean(pending) or cells:
        raise ValueError('Table must end with a complete row')
    return {'kind': 'table', 'caption': caption, 'columns': columns,
            'rows': rows, 'header_rows': header_rows or 1, 'notes': notes}


def parse_figure(node):
    image = caption = None
    for child in node.nodelist:
        if isinstance(child, LatexMacroNode):
            if child.macroname == 'includegraphics':
                image = plain(argument(child))
            elif child.macroname == 'caption':
                caption = clean(inline(argument(child)))
            elif child.macroname not in ('centering', 'label'):
                raise ValueError(f'Unsupported figure command: {child.macroname}')
        elif isinstance(child, LatexCharsNode) and child.chars.strip():
            raise ValueError('Unexpected figure text')
        elif not isinstance(child, (LatexCharsNode, LatexCommentNode)):
            raise ValueError('Unsupported figure content')
    if image is None or caption is None:
        raise ValueError('Each figure requires an image and caption')
    return dict(kind='figure', image=image, caption=caption)


def blocks(text):
    result, pending = [], []

    def flush():
        nonlocal pending
        runs = clean(pending)
        if runs:
            result.append({'kind': 'paragraph', 'runs': runs})
        pending = []

    for node in parse(text):
        if isinstance(node, LatexCharsNode):
            for i, chunk in enumerate(re.split(r'\n\s*\n', node.chars)):
                if i:
                    flush()
                pending.append({'text': re.sub(r'\s+', ' ', chunk)})
        elif isinstance(node, LatexMacroNode) and node.macroname in ('section', 'subsection', 'subsubsection'):
            flush()
            result.append({'kind': node.macroname, 'runs': clean(inline(argument(node)))})
        elif isinstance(node, LatexEnvironmentNode):
            flush()
            if node.environmentname == 'table':
                result.append(parse_table(node))
            elif node.environmentname == 'figure':
                result.append(parse_figure(node))
            elif node.environmentname == 'equation':
                raw = ''.join(n.latex_verbatim() for n in node.nodelist)
                result.append({'kind': 'equation', 'math': re.sub(r'\\label\{[^}]*\}', '', raw).strip()})
            elif node.environmentname in ('enumerate', 'itemize'):
                items, current = [], []
                for child in node.nodelist:
                    if isinstance(child, LatexMacroNode) and child.macroname == 'item':
                        if current:
                            items.append(current)
                        current = []
                    elif not current and isinstance(child, LatexCharsNode) and not child.chars.strip():
                        continue
                    else:
                        current.append(child)
                if current:
                    items.append(current)
                for number, item in enumerate(items, 1):
                    item_blocks = blocks(''.join(n.latex_verbatim() for n in item))
                    for index, block in enumerate(item_blocks):
                        if block['kind'] != 'paragraph':
                            raise ValueError('Only paragraph list items are supported')
                        block['kind'] = 'list_item'
                        block['number'] = number if index == 0 else None
                        block['ordered'] = node.environmentname == 'enumerate'
                        result.append(block)
            else:
                raise ValueError(f'Unsupported environment {node.environmentname}; add explicit handling before export')
        else:
            pending.extend(inline([node]))
    flush()
    return result


def bibliography():
    database = bibtexparser.loads((ROOT / 'references.bib').read_text(encoding='utf-8'))
    references = []
    for entry in database.entries:
        def value(key):
            return plain(parse(entry.get(key, '')))
        names = []
        authors = entry.get('author', '')
        # Corporate authors protect "and" inside braces (e.g. NIST).
        boundaries = [m for m in re.finditer(r'\s+and\s+', authors)
                      if authors[:m.start()].count('{') == authors[:m.start()].count('}')]
        people, start = [], 0
        for match in boundaries:
            people.append(authors[start:match.start()])
            start = match.end()
        people.append(authors[start:])
        for person in people:
            if person.startswith('{') or ',' not in person:
                names.append(plain(parse(person)))
            else:
                last, first = person.split(',', 1)
                initials = ['-'.join(part[0] + '.' for part in word.split('-') if part)
                            for word in first.split()]
                names.append(' '.join(initials) + ' ' + plain(parse(last)))
        fields = [', '.join(names) + ': ' + value('title')]
        for key in ('journal', 'booktitle', 'series', 'institution', 'publisher',
                    'address', 'howpublished'):
            if value(key):
                fields.append(value(key))
        if value('type'):
            fields.append(value('type') + (' ' + value('number') if value('number') else ''))
        else:
            for key, prefix in [('volume', 'Vol. '), ('number', 'No. ')]:
                if value(key):
                    fields.append(prefix + value(key))
        if value('chapter'):
            fields.append('Ch. ' + value('chapter'))
        date = ' '.join(value(k) for k in ('month', 'year') if value(k))
        if date:
            fields.append(date)
        for key, prefix in [('pages', 'pp. '), ('doi', 'doi: '), ('url', ''), ('note', '')]:
            if key == 'url' and value('doi') and value('url') == 'https://doi.org/' + value('doi'):
                continue
            if value(key):
                fields.append(prefix + value(key))
        references.append({'key': entry['ID'], 'text': ', '.join(fields) + '.'})
    if not references:
        raise ValueError('No bibliography entries parsed')
    return references


def read_manuscript():
    wrapper = re.sub(r'(?m)^%.*$', '', (ROOT / 'main.tex').read_text(encoding='utf-8'))
    inputs = re.findall(r'\\input\{([^}]+)\}', wrapper)
    sources = {}

    def expand(name, stack=()):
        path = (ROOT / name).resolve()
        if not path.is_relative_to(ROOT) or name in stack:
            raise ValueError(f'Unsafe or cyclic input: {name}')
        sources[name] = sha256(path.read_bytes()).hexdigest()
        text = path.read_text(encoding='utf-8')
        text = re.sub(r'(?m)^\s*%.*$', '', text)
        return re.sub(r'\\input\{([^}]+)\}',
                      lambda m: expand(m[1], (*stack, name)), text)

    contents = {name: expand(name) for name in inputs}
    meta = {n.macroname: argument(n) for n in parse(contents['content/metadata.tex']) if isinstance(n, LatexMacroNode)}
    notes = [n for n in meta['author'] if isinstance(n, LatexMacroNode) and n.macroname == 'thanks']
    if len(notes) != 1:
        raise ValueError('This template adapter expects one author affiliation footnote')
    data = {'title': plain(meta['title']),
            'author': plain([n for n in meta['author'] if n not in notes]),
            'affiliation': plain(argument(notes[0])),
            'abstract': clean(inline(parse(contents['content/abstract.tex']))),
            'keywords': clean(inline(parse(contents['content/keywords.tex']))),
            'blocks': [], 'references': bibliography(),
            'include_uncited': r'\nocite{*}' in wrapper,
            'source_sha256': sources | {name: sha256((ROOT / name).read_bytes()).hexdigest()
                                       for name in ['main.tex', 'references.bib']}}
    for name, text in contents.items():
        if re.match(r'content/\d', name):
            try:
                data['blocks'].extend(blocks(text))
            except Exception as exc:
                raise ValueError(f'{name}: {exc}') from exc
    if not data['blocks']:
        raise ValueError('No manuscript body found')
    for block in data['blocks']:
        if block['kind'] == 'figure':
            image = (ROOT / block['image']).resolve()
            if not image.is_relative_to(ROOT):
                raise ValueError('Figure must be inside sjee/')
            data['source_sha256'][block['image']] = sha256(image.read_bytes()).hexdigest()
    resolve_links(data, contents)
    return data


def resolve_links(data, contents):
    """Resolve source labels and cite order independently of stale PDF aux files."""
    labels, cited = {}, []
    counters = dict(section=0, subsection=0, subsubsection=0, table=0, equation=0, figure=0)
    label_kinds = {}
    current = None
    current_kind = None

    def visit(nodes):
        nonlocal current, current_kind
        for node in nodes:
            if isinstance(node, LatexMacroNode):
                name = node.macroname
                if name in ('section', 'subsection', 'subsubsection'):
                    counters[name] += 1
                    if name == 'section':
                        counters['subsection'] = counters['subsubsection'] = 0
                    elif name == 'subsection':
                        counters['subsubsection'] = 0
                    levels = ['section', 'subsection', 'subsubsection']
                    current = '.'.join(str(counters[k]) for k in levels[:levels.index(name)+1])
                    current_kind = name
                elif name == 'label':
                    key = plain(argument(node))
                    if key in labels or current is None:
                        raise ValueError(f'Duplicate or unbound label: {key}')
                    labels[key] = current
                    label_kinds[key] = current_kind
                elif name == 'cite':
                    for key in plain(argument(node)).split(','):
                        if key.strip() not in cited:
                            cited.append(key.strip())
            elif isinstance(node, LatexEnvironmentNode):
                saved = current, current_kind
                if node.environmentname in ('table', 'equation', 'figure'):
                    counters[node.environmentname] += 1
                    current = str(counters[node.environmentname])
                    current_kind = node.environmentname
                visit(node.nodelist)
                current, current_kind = saved
            elif isinstance(node, LatexGroupNode):
                visit(node.nodelist)

    for name, text in contents.items():
        if re.match(r'content/(?:\d|abstract|keywords)', name):
            visit(parse(text))
    by_key = {r['key']: r for r in data['references']}
    for key in cited:
        if key not in by_key:
            raise ValueError(f'Unknown citation key: {key}')
    order = cited + ([k for k in by_key if k not in cited]
                     if data.get('include_uncited', True) else [])
    data['references'] = [by_key[k] for k in order]
    numbers = {key: i for i, key in enumerate(order, 1)}

    def resolve(value):
        if isinstance(value, list):
            for item in value:
                resolve(item)
            # Bold the whole table reference, including its preceding word.
            for index in range(len(value) - 1, -1, -1):
                item = value[index]
                if isinstance(item, dict) and item.get('reference_kind') == 'table':
                    item['bold'] = True
                    previous_index = index - 1
                    while previous_index >= 0 and not value[previous_index].get('text', '').strip():
                        value[previous_index]['bold'] = True
                        previous_index -= 1
                    previous = value[previous_index] if previous_index >= 0 else {}
                    match = re.search(r'(Tables?\s*)$', previous.get('text', ''))
                    if match:
                        before = previous['text'][:match.start()]
                        previous['text'] = match[0]
                        previous['bold'] = True
                        if before:
                            value.insert(previous_index, dict(text=before))
        elif isinstance(value, dict):
            if 'link_kind' in value:
                kind, target = value.pop('link_kind'), value.pop('target')
                if kind == 'cite':
                    ids = sorted(set(numbers[k.strip()] for k in target.split(',')))
                    groups = []
                    for n in ids:
                        if groups and n == groups[-1][-1] + 1:
                            groups[-1].append(n)
                        else:
                            groups.append([n])
                    text = ', '.join(f'{g[0]}\u2013{g[-1]}' if len(g) > 2 else ', '.join(map(str,g)) for g in groups)
                    value['text'] = '[' + text + ']'
                else:
                    if target not in labels:
                        raise ValueError(f'Unknown cross-reference label: {target}')
                    value['text'] = '(' + labels[target] + ')' if kind == 'eqref' else labels[target]
                    value['reference_kind'] = label_kinds[target]
            for child in value.values():
                if isinstance(child, (dict, list)):
                    resolve(child)
    resolve(data)
    data['citation_numbers'] = numbers
    data['labels'] = labels


def omml(tex):
    # Keep upright unit names and textual subscripts as single native math runs.
    # Splitting them into separate identifiers makes some readers add letter spacing.
    tex = tex.replace(r'\mbox{-}', '-')
    tex = re.sub(r'\\mathrm\{([A-Za-z0-9-]+)\}', r'\\text{\1}', tex)
    xml = mathml_omml(latex_mathml(tex))
    # mathml2omml 0.0.2 incorrectly closes groupChrPr as groupChr for
    # over-accents such as widetilde. Repair only that known malformed tag.
    xml = re.sub(r'(<m:groupChrPr>[^<]*(?:<(?!/?m:groupChr)[^>]+>[^<]*)*)</m:groupChr>',
                 r'\1</m:groupChrPr>', xml)
    root = etree.fromstring(('<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">' + xml + '</root>').encode('utf-8'))
    # A tilde is an accent, not a grouping character. Word/LibreOffice can
    # omit groupChr tildes even when the repaired XML is well formed.
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    for group in root.xpath('.//m:groupChr[m:groupChrPr/m:chr[@m:val="~"]]', namespaces=ns):
        group.tag = qn('m:acc')
        props = group.find(qn('m:groupChrPr'))
        props.tag = qn('m:accPr')
        props.find(qn('m:chr')).set(qn('m:val'), '\u0303')
        for child in list(props):
            if child.tag != qn('m:chr'):
                props.remove(child)
    # Set full-size math to 11 pt without flattening its native equation structure.
    for run in root.iter(qn('m:r')):
        math_props = run.find(qn('m:rPr'))
        text_node = run.find(qn('m:t'))
        style = math_props.find(qn('m:sty')) if math_props is not None else None
        if (style is not None and style.get(qn('m:val')) == 'p'
                and text_node is not None and len(text_node.text or '') > 1
                and (text_node.text or '').isalpha()):
            if math_props.find(qn('m:nor')) is None:
                math_props.append(OxmlElement('m:nor'))
        props = OxmlElement('w:rPr')
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '22')
        props.append(size)
        run.insert(1 if run.find(qn('m:rPr')) is not None else 0, props)
    return root[0]


def write_runs(paragraph, runs, warnings):
    for item in runs:
        if 'math' in item:
            try:
                paragraph._p.append(omml(item['math']))
            except Exception as exc:
                raise ValueError(f"Native equation conversion failed: {item['math']}") from exc
        else:
            run = paragraph.add_run(item['text'])
            if item.get('bold'):
                run.bold = True
            if item.get('italic'):
                run.italic = True
            if item.get('code'):
                run.font.name = 'Courier New'


def sentence_case(runs):
    """Sentence-case subheadings, preserving acronyms and mixed-case names."""
    runs = deepcopy(runs)
    first = True
    for item in runs:
        if 'text' not in item:
            continue
        def word(match):
            nonlocal first
            token = match[0]
            if first:
                first = False
                return token
            return token.lower() if token.istitle() else token
        item['text'] = re.sub(r'[A-Za-z][A-Za-z0-9]*', word, item['text'])
    return runs


def figure_assets(relative):
    source = ROOT / relative
    target = BUILD / source.stem
    if source.suffix.lower() != '.pdf':
        raise ValueError('Current figure export expects a vector PDF source')
    converter = which('pdftocairo')
    if not converter:
        raise ValueError('pdftocairo is required for vector figures (included with TeX Live).')
    svg = target.with_suffix('.svg')
    png = target.with_suffix('.png')
    subprocess.run([converter, '-svg', str(source), str(svg)], check=True,
                   capture_output=True)
    subprocess.run([converter, '-png', '-singlefile', '-scale-to', '2400',
                    str(source), str(target)], check=True, capture_output=True)
    return svg, png


def insert_figure(document, paragraph, relative, number):
    svg, png = figure_assets(relative)
    shape = paragraph.add_run().add_picture(str(png), width=Cm(12.5))
    # Modern Word renders the SVG; older readers retain a high-resolution fallback.
    svg_part = Part(PackURI(f'/word/media/figure{number}.svg'), 'image/svg+xml',
                    svg.read_bytes(), document.part.package)
    rid = document.part.relate_to(svg_part, RT.IMAGE)
    blip = shape._inline.find('.//' + qn('a:blip'))
    extensions = OxmlElement('a:extLst')
    extension = OxmlElement('a:ext')
    extension.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    svg_blip = etree.Element(
        '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
        nsmap={'asvg': 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
    svg_blip.set(qn('r:embed'), rid)
    extension.append(svg_blip)
    extensions.append(extension)
    blip.append(extensions)


def list_numbering(document, ordered):
    numbering = document.part.numbering_part.element
    style = document.styles['JText-numberedList' if ordered else 'JText-unorderedList']
    base_id = style.element.find('.//' + qn('w:numId')).get(qn('w:val'))
    base = next(n for n in numbering.findall(qn('w:num'))
                if n.get(qn('w:numId')) == base_id)
    abstract = int(base.find(qn('w:abstractNumId')).get(qn('w:val')))
    number = numbering.add_num(abstract)
    number.add_lvlOverride(ilvl=0).add_startOverride(1)
    return number.numId


def fill_footnote(document, text):
    part = next(p for p in document.part.package.parts if str(p.partname) == '/word/footnotes.xml')
    root = parse_xml(part.blob)
    for note in list(root):
        if note.get(qn('w:id')) == '2':
            root.remove(note)
        elif note.get(qn('w:id')) == '1':
            node = note.find(qn('w:p'))
            marker = deepcopy(node.find('.//' + qn('w:footnoteRef')).getparent())
            paragraph = Paragraph(node, document._body)
            paragraph.clear()
            node.append(marker)
            paragraph.add_run(' ' + text)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def render(data, temporary):
    copyfile(SOURCE, temporary)
    document = Document(temporary)
    prototypes = {}
    for p in document.paragraphs:
        prototypes.setdefault(p.style.style_id, deepcopy(p._p))
    marker = deepcopy(prototypes['JAuthor'].find('.//' + qn('w:footnoteReference')).getparent())
    table_properties = deepcopy(document.tables[0]._tbl.tblPr)
    body = document._element.body
    for node in list(body):
        if node.tag != qn('w:sectPr'):
            body.remove(node)
    warnings = []

    def paragraph(style, runs=None):
        node = deepcopy(prototypes[style])
        p = Paragraph(node, document._body)
        p.clear()
        body.insert(len(body) - 1, node)
        if runs:
            write_runs(p, runs, warnings)
        return p

    title = data['title']
    if '\n' not in title and len(title) > 48:
        title = '\n'.join(textwrap.wrap(title, width=45, break_long_words=False))
    paragraph('JTitle', [{'text': title}])
    paragraph('JAuthor', [{'text': data['author']}])._p.append(marker)
    paragraph('JAbstract', [{'text': 'Abstract: ', 'bold': True}] + data['abstract'])
    keywords = deepcopy(data['keywords'])
    for item in keywords:
        if 'text' in item:
            item['text'] = re.sub(r'(^|,\s*)([a-z])',
                                  lambda m: m[1] + m[2].upper(), item['text'])
    paragraph('JAbstract', [{'text': 'Keywords: ', 'bold': True}] + keywords)
    fill_footnote(document, data['affiliation'])
    for part in document.part.package.parts:
        if str(part.partname).startswith('/word/header') and str(part.partname).endswith('.xml'):
            for node in part.element.findall('.//' + qn('w:p')):
                Paragraph(node, document._body).clear()
    section = subsection = subsubsection = table_number = equation = figure = 0
    list_id = None
    after_equation = False
    for block in data['blocks']:
        kind = block['kind']
        if kind == 'section':
            section += 1
            subsection = 0
            subsubsection = 0
            paragraph('JSubtitle', [{'text': f'{section}\t'}] + block['runs'])
        elif kind == 'subsection':
            subsection += 1
            subsubsection = 0
            paragraph('JSubSubTitle', [{'text': f'{section}.{subsection}\t'}] + sentence_case(block['runs']))
        elif kind == 'subsubsection':
            subsubsection += 1
            paragraph('JSubSubTitle', [{'text': f'{section}.{subsection}.{subsubsection}\t'}] + sentence_case(block['runs']))
        elif kind == 'list_item':
            ordered = block.get('ordered', True)
            if block['number'] == 1:
                list_id = list_numbering(document, ordered)
            style = 'JText-numberedList' if ordered else 'JText-unorderedList'
            p = paragraph(style if block['number'] is not None else 'JText-noIndent', block['runs'])
            if block['number'] is not None:
                props = p._p.get_or_add_pPr().get_or_add_numPr()
                props.get_or_add_ilvl().val = 0
                props.get_or_add_numId().val = list_id
            else:
                p.paragraph_format.left_indent = Cm(.7)
        elif kind == 'paragraph':
            paragraph('JText-noIndent' if after_equation else 'JText', block['runs'])
        elif kind == 'equation':
            equation += 1
            p = paragraph('JEquation')
            p.add_run('\t')
            write_runs(p, [{'math': block['math']}], warnings)
            p.add_run(f'\t({equation})')
        elif kind == 'figure':
            figure += 1
            p = paragraph('JFigure')
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            insert_figure(document, p, block['image'], figure)
            p = paragraph('JFigure', [{'text': f'Fig. {figure} \u2013 '}] + block['caption'])
            p.paragraph_format.keep_together = True
        elif kind == 'table':
            table_number += 1
            p = paragraph('JFigure', [{'text': f'Table {table_number}', 'bold': True}, {'text': '\n'}] +
                          [dict(r, italic=True) for r in block['caption']])
            p.paragraph_format.keep_with_next = True
            ncols = len(block['columns'])
            table = document.add_table(rows=len(block['rows']), cols=ncols)
            props = deepcopy(table_properties)
            # Original tblStyle points at a paragraph style; retain explicit
            # template grid borders/alignment, not that invalid style reference.
            old_style = props.find(qn('w:tblStyle'))
            if old_style is not None:
                props.remove(old_style)
            props.find(qn('w:tblW')).set(qn('w:w'), '7088')
            table._tbl.replace(table._tbl.tblPr, props)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            weights = [5.0] * ncols
            for row in block['rows']:
                col = 0
                for cell in row:
                    length = len(''.join(r.get('text', r.get('math', '')) for r in cell['runs']))
                    if cell['span'] == 1:
                        weights[col] = max(weights[col], min(length, 23) + 2)
                    col += cell['span']
            widths = [Cm(12.5 * w / sum(weights)) for w in weights]
            for i, column in enumerate(table.columns):
                column.width = widths[i]
            for ri, row in enumerate(block['rows']):
                word_row = table.rows[ri]
                trpr = word_row._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if ri < block['header_rows']:
                    trpr.append(OxmlElement('w:tblHeader'))
                for i, cell in enumerate(word_row.cells):
                    cell.width = widths[i]
                col = 0
                for cell in row:
                    target = word_row.cells[col]
                    if cell['span'] > 1:
                        target = target.merge(word_row.cells[col + cell['span'] - 1])
                    target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = target.paragraphs[0]
                    p.style = document.styles['JTableText']
                    # Keep a short table together, moving it as a unit when needed.
                    p.paragraph_format.keep_with_next = ri < len(block['rows']) - 1
                    p.paragraph_format.keep_together = True
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if block['columns'][col] == 'l' else (
                        WD_ALIGN_PARAGRAPH.RIGHT if block['columns'][col] == 'r'
                        else WD_ALIGN_PARAGRAPH.CENTER)
                    write_runs(p, cell['runs'], warnings)
                    if ri == 0 and any('GPU throughput at' in r.get('text', '')
                                       for r in block['caption']) and col >= 2:
                        p.add_run(' [MB/s]')
                    col += cell['span']
            if block.get('notes'):
                note = paragraph('JText-noIndent', block['notes'])
                for run in note.runs:
                    run.font.size = Pt(9)
        else:
            raise ValueError(f'Unhandled block kind {kind}')
        after_equation = kind == 'equation'
    paragraph('JSubtitle', [{'text': f'{section + 1}\tReferences'}])
    for reference in data['references']:
        p = paragraph('JReference', [{'text': reference['text']}])
        # Long documentation URLs should not stretch short preceding lines.
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.core_properties.title = data['title']
    document.core_properties.author = data['author']
    document.core_properties.subject = 'SJEE manuscript draft'
    document.core_properties.comments = 'Generated from content/*.tex and references.bib; edit those sources.'
    used = {v for el in body.iter() for k, v in el.attrib.items() if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
    for rid, rel in list(document.part.rels.items()):
        if rid not in used and rel.reltype.rsplit('/', 1)[-1] in ('image', 'oleObject'):
            document.part.drop_rel(rid)
    document.save(temporary)
    return warnings


def validate(path, data):
    with ZipFile(SOURCE) as source, ZipFile(path) as result:
        for name in ['word/styles.xml']:
            if etree.tostring(parse_xml(source.read(name)), method='c14n') != etree.tostring(parse_xml(result.read(name)), method='c14n'):
                raise ValueError(f'Template styles/numbering changed: {name}')
        original_numbers = parse_xml(source.read('word/numbering.xml'))
        output_numbers = parse_xml(result.read('word/numbering.xml'))
        for original in original_numbers:
            identifier = original.get(qn('w:numId')) or original.get(qn('w:abstractNumId'))
            candidates = [n for n in output_numbers if n.tag == original.tag
                          and (n.get(qn('w:numId')) or n.get(qn('w:abstractNumId'))) == identifier]
            if len(candidates) != 1 or etree.tostring(original, method='c14n') != etree.tostring(candidates[0], method='c14n'):
                raise ValueError('Original list definition changed')
        src = parse_xml(source.read('word/document.xml'))
        out = parse_xml(result.read('word/document.xml'))
        for tag in ('pgSz', 'pgMar', 'cols'):
            if src.find('.//' + qn('w:' + tag)).attrib != out.find('.//' + qn('w:' + tag)).attrib:
                raise ValueError(f'Page setting changed: {tag}')
        if len(out.findall('.//' + qn('w:tbl'))) != sum(b['kind'] == 'table' for b in data['blocks']):
            raise ValueError('Table count mismatch')
        for name in result.namelist():
            if name.startswith('word/header') and name.endswith('.xml'):
                if ''.join(parse_xml(result.read(name)).itertext()).strip():
                    raise ValueError(f'Sample header text remains: {name}')
        if out.findall('.//' + qn('w:object')):
            raise ValueError('Legacy template equation objects remain')
        expected_figures = sum(b['kind'] == 'figure' for b in data['blocks'])
        if len(out.findall('.//' + qn('wp:inline'))) != expected_figures:
            raise ValueError('Figure count mismatch')
        if len([n for n in result.namelist() if n.startswith('word/media/')
                and n.endswith('.svg')]) != expected_figures:
            raise ValueError('Missing vector figure')
        def math_count(value):
            if isinstance(value, dict):
                return int('math' in value) + sum(math_count(v) for v in value.values())
            if isinstance(value, list):
                return sum(math_count(v) for v in value)
            return 0
        if len(out.findall('.//' + qn('m:oMath'))) != math_count(data):
            raise ValueError('Native equation count mismatch')
        if '[MATH CLEANUP:' in ''.join(out.itertext()):
            raise ValueError('Unconverted math remains')


def main():
    if sha256(SOURCE.read_bytes()).hexdigest() != TEMPLATE_HASH:
        raise ValueError('Source template changed; inspect and adapt the template mapping first')
    if OUTPUT.is_symlink() or (OUTPUT.exists() and SOURCE.samefile(OUTPUT)):
        raise ValueError('Output must not link to the original template')
    data = read_manuscript()
    BUILD.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='sjee-docx-', dir=BUILD) as directory:
        temporary = Path(directory) / 'main_sjee.docx'
        warnings = render(data, temporary)
        validate(temporary, data)
        for name, digest in data['source_sha256'].items():
            if sha256((ROOT / name).read_bytes()).hexdigest() != digest:
                raise ValueError(f'Source changed during generation: {name}; rerun')
        if sha256(SOURCE.read_bytes()).hexdigest() != TEMPLATE_HASH:
            raise ValueError('Original template hash changed')
        temporary.replace(OUTPUT)
    (BUILD / 'manuscript.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    report = {'output': str(OUTPUT), 'generator': 'SJEE template copy; no Pandoc',
              'template_sha256': TEMPLATE_HASH, 'source_sha256': data['source_sha256'],
              'tables': sum(b['kind'] == 'table' for b in data['blocks']),
              'equations': sum(b['kind'] == 'equation' for b in data['blocks']),
              'figures': sum(b['kind'] == 'figure' for b in data['blocks']),
              'references': len(data['references']), 'warnings': warnings,
              'manual_review': ['Review resolved numeric citations and cross-references before submission.',
                                'Review equation appearance, table pagination and reference punctuation in Word.',
                                'Caption/equation numbers are generated literals, not Word cross-reference fields.']}
    (BUILD / 'build_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Created {OUTPUT}: {report["tables"]} tables, {report["figures"]} figures, {report["equations"]} display equation(s), {report["references"]} references.')
    for warning in warnings:
        print('WARNING:', warning)
    print('Preserved SJEE styles/page settings. Cleared sample headers. Original template unchanged.')
    print('Source: content/*.tex and references.bib. Report: build/docx/build_report.json')


if __name__ == '__main__':
    main()
